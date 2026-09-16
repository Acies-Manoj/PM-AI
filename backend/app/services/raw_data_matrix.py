"""Parsing for the Raw Data Explorer step: an aggregated trip summary plus
one or two "data point matrix" files (one column per trip, header = Serial
Number / Trip ID / unit, rows = sequential raw readings at a fixed interval,
no timestamps). See raw-data-explorer/app.py for the standalone prototype
this was ported from -- same parsing rules, verified against the real demo
files (Merged_Temperature_Light_2024-2025.xlsx / Data Point Extraction
Matrix.xlsx)."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

import openpyxl
import pandas as pd

AGGREGATED_SHEET_HINTS = ["temperature", "light", "trip", "summary"]
MATRIX_SHEET_HINTS = ["extraction matrix", "matrix", "temperature", "temp", "light"]

SERIAL_COLUMN_OPTIONS = ["Serial Number", "Serial No", "Device Serial Number"]
TRIP_COLUMN_OPTIONS = ["Trip ID", "TripID", "Trip Id"]
START_COLUMN_OPTIONS = [
    "Segment Start Date Time",
    "Trip Start",
    "Start Date Time",
    "Actual Departure Time CET",
]
LABEL_EXTRA_COLUMNS = ["Product", "Origin", "Destination", "Container Name"]


def list_sheets(file_bytes: bytes) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def guess_sheet(sheet_names: list[str], prefer: list[str]) -> str:
    lowered = [s.lower() for s in sheet_names]
    for keyword in prefer:
        for i, name in enumerate(lowered):
            if keyword in name:
                return sheet_names[i]
    return sheet_names[0]


@dataclass
class AggregatedData:
    df: pd.DataFrame
    serial_col: str
    trip_col: str
    start_col: str
    label_parts: list[str]


_HEADER_HINT_TOKENS = [
    opt.lower() for opt in (SERIAL_COLUMN_OPTIONS + TRIP_COLUMN_OPTIONS + START_COLUMN_OPTIONS)
]


def _detect_header_row(file_bytes: bytes, sheet_name: str, max_scan_rows: int = 15) -> int:
    """Finds the row that actually holds the column headers, for exports
    that put a title row (and often a blank row) above them -- e.g. a
    "Tabular Shipment Data" export whose row 1 is just the report title,
    with the real headers on row 3. Returns a 0-based row index for
    pandas' `header=`.

    Requires at least 2 of the known serial/trip/start column names to
    appear in a row before calling it the header -- one coincidental match
    (a data cell that happens to read "Trip Start") isn't enough to be
    confident, but a real header row always carries several at once. Falls
    back to row 0 (pandas' own default) if nothing scores higher, so a file
    whose header already sits on row 1 behaves exactly as before."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True)):
            cells = {re.sub(r"\s+", " ", str(c)).strip().lower() for c in row if c is not None}
            if sum(1 for token in _HEADER_HINT_TOKENS if token in cells) >= 2:
                return i
        return 0
    finally:
        wb.close()


def parse_aggregated(file_bytes: bytes, sheet_name: str) -> AggregatedData:
    header_row = _detect_header_row(file_bytes, sheet_name)
    df = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_name, header=header_row)

    def find_col(options: list[str]) -> str | None:
        for opt in options:
            for c in df.columns:
                if str(c).strip().lower() == opt.lower():
                    return c
        return None

    serial_col = find_col(SERIAL_COLUMN_OPTIONS)
    trip_col = find_col(TRIP_COLUMN_OPTIONS)
    start_col = find_col(START_COLUMN_OPTIONS)

    if start_col is None:
        dt_cols = [c for c in df.columns if pd.api.types.is_datetime64_any_dtype(df[c])]
        start_candidates = [c for c in dt_cols if "start" in str(c).lower()]
        start_col = start_candidates[0] if start_candidates else (dt_cols[0] if dt_cols else None)
    elif not pd.api.types.is_datetime64_any_dtype(df[start_col]):
        df[start_col] = pd.to_datetime(df[start_col], errors="coerce")

    if serial_col is None or trip_col is None or start_col is None:
        missing = [
            name
            for name, col in [("Serial Number", serial_col), ("Trip ID", trip_col), ("start timestamp", start_col)]
            if col is None
        ]
        raise ValueError(f"Could not find the following required column(s) in the aggregated data: {', '.join(missing)}.")

    label_parts = [trip_col, serial_col]
    for extra in LABEL_EXTRA_COLUMNS:
        col = find_col([extra])
        if col is not None:
            label_parts.append(col)

    return AggregatedData(df=df, serial_col=serial_col, trip_col=trip_col, start_col=start_col, label_parts=label_parts)


@dataclass
class MatrixData:
    # (serial, trip_id) -> column index into each row tuple in data_rows
    columns: dict[tuple[str, int], int]
    data_rows: list[tuple]
    unit: str


def parse_matrix(file_bytes: bytes, sheet_name: str) -> MatrixData:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]

        header_row_idx = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            first = str(row[0]).strip().lower() if row and row[0] else ""
            if first == "point number":
                header_row_idx = i
                break
        if header_row_idx is None:
            header_row_idx = 1

        header = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))

        columns: dict[tuple[str, int], int] = {}
        unit = ""
        for idx, h in enumerate(header):
            if idx == 0 or not h:
                continue
            parts = [p.strip() for p in str(h).split("\n") if p.strip() != ""]
            if len(parts) < 2:
                continue
            serial = parts[0]
            try:
                trip_id = int(float(parts[1]))
            except ValueError:
                continue
            if len(parts) >= 3:
                unit = parts[2]
            columns[(serial, trip_id)] = idx

        data_rows = list(ws.iter_rows(min_row=header_row_idx + 1, values_only=True))
        return MatrixData(columns=columns, data_rows=data_rows, unit=unit)
    finally:
        wb.close()


@dataclass
class TripRow:
    serial: str
    trip_id: int
    label: str
    start_time: pd.Timestamp


def build_trip_rows(agg: AggregatedData) -> list[TripRow]:
    """One row per trip with a usable (Serial Number, Trip ID, start
    timestamp) -- rows missing any of those three are dropped, since they
    can't be joined against a matrix or anchor a timestamp. Sorted by start
    time, matching the standalone prototype's trip picker order.

    Some exports carry more than one aggregated row per trip -- e.g. a
    separate row per sensor channel (Temperature/Light), or per product on a
    multi-product trip -- so this also de-duplicates by (serial, trip_id),
    keeping the first (earliest-sorted) row, rather than showing the same
    trip 2-4 times in the picker."""
    df = agg.df.dropna(subset=[agg.serial_col, agg.trip_col, agg.start_col]).copy()
    df = df.sort_values(agg.start_col)

    rows: list[TripRow] = []
    seen: set[tuple[str, int]] = set()
    for _, row in df.iterrows():
        try:
            trip_id = int(float(row[agg.trip_col]))
        except (TypeError, ValueError):
            continue
        serial = str(row[agg.serial_col])
        key = (serial, trip_id)
        if key in seen:
            continue
        seen.add(key)
        label = " | ".join(str(row[c]) for c in agg.label_parts if c in row and pd.notna(row[c]))
        rows.append(TripRow(serial=serial, trip_id=trip_id, label=label, start_time=row[agg.start_col]))
    return rows


def get_series(matrix: MatrixData, serial: str, trip_id: int) -> list[float | None] | None:
    """None => trip not present in this matrix at all. [] => present but no
    readings. Otherwise the raw values in point order, trimmed of the
    trailing padding (every trip column is padded to the sheet's max row
    count with blanks past its own last real reading)."""
    col_idx = matrix.columns.get((serial, trip_id))
    if col_idx is None:
        return None
    raw = [row[col_idx] if col_idx < len(row) else None for row in matrix.data_rows]
    last_idx = None
    for i, v in enumerate(raw):
        if v is not None:
            last_idx = i
    if last_idx is None:
        return []
    return list(raw[: last_idx + 1])
