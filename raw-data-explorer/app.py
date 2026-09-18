"""Raw Cold-Chain Data Explorer.

Three-upload tool:
  1. Aggregated Data Upload   - one row per trip, incl. Serial Number, Trip ID,
                                 and a segment/trip start timestamp.
  2. Data Point Matrix (Temperature) - one column per trip (header = Serial
                                 Number / Trip ID / unit), rows are sequential
                                 raw readings at a fixed interval.
  3. Data Point Matrix (Light)       - same layout as above, for light readings.

Pick a trip -> its start timestamp from the aggregated file is used as the
timestamp of raw point #1, later points are spaced by the reading interval.
"""

from __future__ import annotations

import hashlib
from io import BytesIO

import openpyxl
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

st.set_page_config(page_title="Raw Cold-Chain Data Explorer", layout="wide")

st.title("Raw Cold-Chain Data Explorer")
st.caption(
    "Upload the aggregated trip summary plus the raw data-point matrices for "
    "temperature and light. Search/select a trip below - its raw sensor curve "
    "is reconstructed starting at that trip's own start timestamp."
)


# --------------------------------------------------------------------------- #
# Parsing helpers
# --------------------------------------------------------------------------- #

def _hash(file) -> str:
    return hashlib.sha256(file.getvalue()).hexdigest()


@st.cache_data(show_spinner=False)
def list_sheets(file_bytes: bytes, _key: str) -> list[str]:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def _guess_sheet(sheet_names: list[str], prefer: list[str]) -> str:
    lowered = [s.lower() for s in sheet_names]
    for keyword in prefer:
        for i, name in enumerate(lowered):
            if keyword in name:
                return sheet_names[i]
    return sheet_names[0]


@st.cache_data(show_spinner=False)
def parse_aggregated(file_bytes: bytes, _key: str, sheet_name: str):
    df = pd.read_excel(BytesIO(file_bytes), sheet_name=sheet_name)

    def find_col(options: list[str]):
        for opt in options:
            for c in df.columns:
                if str(c).strip().lower() == opt.lower():
                    return c
        return None

    serial_col = find_col(["Serial Number", "Serial No", "Device Serial Number"])
    trip_col = find_col(["Trip ID", "TripID", "Trip Id"])
    start_col = find_col(
        ["Segment Start Date Time", "Trip Start", "Start Date Time", "Actual Departure Time CET"]
    )

    if start_col is None:
        dt_cols = [c for c in df.columns if pd.api.types.is_datetime64_any_dtype(df[c])]
        start_candidates = [c for c in dt_cols if "start" in str(c).lower()]
        start_col = start_candidates[0] if start_candidates else (dt_cols[0] if dt_cols else None)
    elif not pd.api.types.is_datetime64_any_dtype(df[start_col]):
        df[start_col] = pd.to_datetime(df[start_col], errors="coerce")

    label_parts = [c for c in [trip_col, serial_col] if c is not None]
    for extra in ["Product", "Origin", "Destination", "Container Name"]:
        col = find_col([extra])
        if col is not None:
            label_parts.append(col)

    return df, serial_col, trip_col, start_col, label_parts


@st.cache_data(show_spinner=False)
def parse_matrix(file_bytes: bytes, _key: str, sheet_name: str):
    wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[sheet_name]

    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        first = str(row[0]).strip().lower() if row and row[0] else ""
        if first == "point number":
            header_row_idx = i
            break
    if header_row_idx is None:
        header_row_idx = 1

    header = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))

    columns: dict[tuple[str, int], int] = {}
    unit_label = ""
    for idx, h in enumerate(header):
        if idx == 0 or not h:
            continue
        parts = [p.strip() for p in str(h).split("\n") if p.strip() != ""]
        if len(parts) < 2:
            continue
        serial = parts[0]
        try:
            trip_id = int(float(parts[1]))
        except ValueError:
            continue
        if len(parts) >= 3:
            unit_label = parts[2]
        columns[(serial, trip_id)] = idx

    data_rows = [
        row for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True)
    ]
    wb.close()
    return columns, data_rows, unit_label


def get_series(columns, data_rows, serial: str, trip_id: int):
    col_idx = columns.get((serial, trip_id))
    if col_idx is None:
        return None
    raw = [row[col_idx] if col_idx < len(row) else None for row in data_rows]
    last_idx = None
    for i, v in enumerate(raw):
        if v is not None:
            last_idx = i
    if last_idx is None:
        return []
    return raw[: last_idx + 1]


# --------------------------------------------------------------------------- #
# Uploads
# --------------------------------------------------------------------------- #

col1, col2, col3 = st.columns(3)
with col1:
    agg_file = st.file_uploader("1. Aggregated Data Upload", type=["xlsx"], key="agg")
with col2:
    temp_file = st.file_uploader("2. Data Point Matrix - Temperature", type=["xlsx"], key="temp")
with col3:
    light_file = st.file_uploader("3. Data Point Matrix - Light", type=["xlsx"], key="light")

if not agg_file:
    st.info("Upload the Aggregated Data file to get started.")
    st.stop()

agg_bytes = agg_file.getvalue()
agg_key = _hash(agg_file)
agg_sheets = list_sheets(agg_bytes, agg_key)
agg_sheet = st.selectbox(
    "Aggregated data sheet", agg_sheets,
    index=agg_sheets.index(_guess_sheet(agg_sheets, ["temperature", "light", "trip", "summary"])),
)
df, serial_col, trip_col, start_col, label_parts = parse_aggregated(agg_bytes, agg_key, agg_sheet)

missing = [name for name, col in [("Serial Number", serial_col), ("Trip ID", trip_col), ("start timestamp", start_col)] if col is None]
if missing:
    st.error(f"Could not find the following required column(s) in the aggregated data: {', '.join(missing)}.")
    st.stop()

temp_data = None
if temp_file:
    temp_bytes = temp_file.getvalue()
    temp_key = _hash(temp_file)
    temp_sheets = list_sheets(temp_bytes, temp_key)
    temp_sheet = st.selectbox(
        "Temperature matrix sheet", temp_sheets,
        index=temp_sheets.index(_guess_sheet(temp_sheets, ["extraction matrix", "matrix", "temperature", "temp"])),
    )
    temp_data = parse_matrix(temp_bytes, temp_key, temp_sheet)

light_data = None
if light_file:
    light_bytes = light_file.getvalue()
    light_key = _hash(light_file)
    light_sheets = list_sheets(light_bytes, light_key)
    light_sheet = st.selectbox(
        "Light matrix sheet", light_sheets,
        index=light_sheets.index(_guess_sheet(light_sheets, ["extraction matrix", "matrix", "light"])),
    )
    light_data = parse_matrix(light_bytes, light_key, light_sheet)

if temp_data is None and light_data is None:
    st.warning("Upload at least one data-point matrix (temperature and/or light) to plot raw readings.")

interval_minutes = st.number_input(
    "Reading interval between raw points (minutes)", min_value=1, value=15, step=1
)

# --------------------------------------------------------------------------- #
# Diagnostics: how much overlap exists between aggregated trips and matrix trips
# --------------------------------------------------------------------------- #

agg_trip_keys = set(zip(df[serial_col].astype(str), pd.to_numeric(df[trip_col], errors="coerce").dropna().astype(int)))
diag_cols = st.columns(3)
diag_cols[0].metric("Trips in aggregated data", len(agg_trip_keys))
if temp_data is not None:
    overlap_temp = len(agg_trip_keys & set(temp_data[0].keys()))
    diag_cols[1].metric("Temperature matrix trips matched", f"{overlap_temp} / {len(temp_data[0])}")
if light_data is not None:
    overlap_light = len(agg_trip_keys & set(light_data[0].keys()))
    diag_cols[2].metric("Light matrix trips matched", f"{overlap_light} / {len(light_data[0])}")

# --------------------------------------------------------------------------- #
# Trip filter / selection
# --------------------------------------------------------------------------- #

def build_label(row) -> str:
    return " | ".join(str(row[c]) for c in label_parts if c in row and pd.notna(row[c]))

df = df.dropna(subset=[serial_col, trip_col, start_col]).copy()
df["_label"] = df.apply(build_label, axis=1)
df = df.sort_values(start_col)

search = st.text_input("Filter trips (matches Trip ID, Serial Number, Product, etc.)", "")
filtered = df[df["_label"].str.contains(search, case=False, na=False)] if search else df

if filtered.empty:
    st.warning("No trips match that filter.")
    st.stop()

selected_label = st.selectbox("Select a trip", filtered["_label"].tolist())
selected_row = filtered[filtered["_label"] == selected_label].iloc[0]

serial = str(selected_row[serial_col])
trip_id = int(selected_row[trip_col])
start_ts = pd.to_datetime(selected_row[start_col])

st.write(f"**Trip start timestamp (point #1):** {start_ts}")

# --------------------------------------------------------------------------- #
# Build raw series & chart
# --------------------------------------------------------------------------- #

fig = go.Figure()
has_any_series = False

if temp_data is not None:
    series = get_series(temp_data[0], temp_data[1], serial, trip_id)
    if series is None:
        st.warning(f"No temperature raw data found for Trip {trip_id} / {serial} in the uploaded temperature matrix.")
    elif series:
        timestamps = [start_ts + pd.Timedelta(minutes=interval_minutes * i) for i in range(len(series))]
        fig.add_trace(go.Scatter(
            x=timestamps, y=series, mode="lines+markers", name=f"Temperature ({temp_data[2] or 'raw'})",
            yaxis="y1",
        ))
        has_any_series = True

if light_data is not None:
    series = get_series(light_data[0], light_data[1], serial, trip_id)
    if series is None:
        st.warning(f"No light raw data found for Trip {trip_id} / {serial} in the uploaded light matrix.")
    elif series:
        timestamps = [start_ts + pd.Timedelta(minutes=interval_minutes * i) for i in range(len(series))]
        fig.add_trace(go.Scatter(
            x=timestamps, y=series, mode="lines+markers", name=f"Light ({light_data[2] or 'raw'})",
            yaxis="y2",
        ))
        has_any_series = True

if has_any_series:
    fig.update_layout(
        title=f"Raw readings - Trip {trip_id} ({serial})",
        xaxis_title="Timestamp",
        yaxis=dict(title="Temperature"),
        yaxis2=dict(title="Light", overlaying="y", side="right"),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
        height=550,
    )
    st.plotly_chart(fig, use_container_width=True)
else:
    st.info("No raw data points were found for this trip in the uploaded matrix file(s).")
